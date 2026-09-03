import pandas as pd
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# 설정
# ============================================================

# analysis 폴더
ANALYSIS_DIR = Path(
    r"D:\VSC0DE_파이썬\prj_etf_validation\us\data\us_etf\analysis"
)

# 분석할 ETF
TICKERS = [
    "TQQQ",    
    "QLD",
    "QQQ",
    "SPY",
    "SCHD",
    "DIA",
    "DDM",
    "UDOW",
]

# 결과 파일
OUTPUT_FILE = ANALYSIS_DIR / "ETF_고점대비하락률_비교.xlsx"

# CSV 결과 파일
CSV_FILE = ANALYSIS_DIR / "ETF_고점대비하락률_비교.csv"


# ============================================================
# ETF별 CSV 읽기
# ============================================================

data_list = []

for ticker in TICKERS:

    file_path = ANALYSIS_DIR / f"{ticker}_analysis.csv"

    if not file_path.exists():
        print(f"[경고] 파일 없음: {file_path}")
        continue

    print(f"[읽는 중] {ticker}")

    try:
        df = pd.read_csv(
            file_path,
            encoding="utf-8-sig"
        )
    except UnicodeDecodeError:
        df = pd.read_csv(
            file_path,
            encoding="cp949"
        )

    # 필수 컬럼 확인
    required_columns = [
        "날짜",
        "고점_대비_하락률"
    ]

    for column in required_columns:

        if column not in df.columns:
            raise ValueError(
                f"{ticker}_analysis.csv에 "
                f"'{column}' 컬럼이 없습니다."
            )

    # 날짜 변환
    df["날짜"] = pd.to_datetime(
        df["날짜"],
        errors="coerce"
    )

    # 잘못된 날짜 제거
    df = df.dropna(
        subset=["날짜"]
    )

    # 필요한 데이터만 추출
    temp = df[
        [
            "날짜",
            "고점_대비_하락률"
        ]
    ].copy()

    # 숫자 변환
    temp["고점_대비_하락률"] = pd.to_numeric(
        temp["고점_대비_하락률"],
        errors="coerce"
    )

    # ETF 이름으로 컬럼 변경
    temp.rename(
        columns={
            "고점_대비_하락률": ticker
        },
        inplace=True
    )

    # 중복 날짜 제거
    temp = temp.drop_duplicates(
        subset=["날짜"]
    )

    data_list.append(temp)


# ============================================================
# 데이터 확인
# ============================================================

if not data_list:
    raise ValueError(
        "읽을 수 있는 ETF analysis CSV 파일이 없습니다."
    )


# ============================================================
# 날짜 기준 병합
# ============================================================

result = data_list[0]

for df in data_list[1:]:

    result = pd.merge(
        result,
        df,
        on="날짜",
        how="outer"
    )


# ============================================================
# 날짜순 정렬
# ============================================================

result = result.sort_values(
    by="날짜"
).reset_index(
    drop=True
)


# ============================================================
# 숫자 반올림
# ============================================================

for ticker in TICKERS:

    if ticker in result.columns:

        result[ticker] = pd.to_numeric(
            result[ticker],
            errors="coerce"
        )

        result[ticker] = result[ticker].round(2)


# ============================================================
# CSV 저장
# ============================================================

csv_result = result.copy()

csv_result["날짜"] = csv_result[
    "날짜"
].dt.strftime("%Y-%m-%d")

csv_result.to_csv(
    CSV_FILE,
    index=False,
    encoding="utf-8-sig"
)

# ============================================================
# Excel 저장
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl",
    date_format="yyyy-mm-dd"
) as writer:

    # ========================================================
    # 전체 Sheet
    # ========================================================

    result.to_excel(
        writer,
        sheet_name="전체",
        index=False
    )

    # ========================================================
    # 연도별 Sheet
    # ========================================================

    years = sorted(
        result["날짜"]
        .dt.year
        .dropna()
        .unique()
    )

    for year in years:

        # 해당 연도 데이터
        yearly_data = result[
            result["날짜"].dt.year == year
        ].copy()

        yearly_data = yearly_data.sort_values(
            by="날짜"
        )

        yearly_data.to_excel(
            writer,
            sheet_name=str(year),
            index=False
        )


# ============================================================
# Excel 서식 적용
# ============================================================

wb = load_workbook(
    OUTPUT_FILE
)


# ============================================================
# 색상 설정
# ============================================================

# 최저값 빨간색 바탕
MIN_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)

# 최저값 글자
MIN_FONT = Font(
    color="9C0006",
    bold=True
)

# 요약행 바탕
SUMMARY_LABEL_FILL = PatternFill(
    fill_type="solid",
    fgColor="E7E6E6"
)

# 헤더
HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

HEADER_FONT = Font(
    bold=True
)


# ============================================================
# Sheet 서식 + 최저값 행 생성
# ============================================================

def format_sheet(ws):

    # --------------------------------------------------------
    # 헤더
    # --------------------------------------------------------

    for cell in ws[1]:

        cell.fill = HEADER_FILL

        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # --------------------------------------------------------
    # 헤더 위치 찾기
    # --------------------------------------------------------

    header_map = {}

    for cell in ws[1]:

        header_map[cell.value] = cell.column

    # --------------------------------------------------------
    # 날짜 컬럼
    # --------------------------------------------------------

    date_col = header_map.get("날짜")

    if date_col:

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(
                row=row,
                column=date_col
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.number_format = "yyyy-mm-dd"

    # --------------------------------------------------------
    # 현재 데이터 마지막 행
    # --------------------------------------------------------

    last_data_row = ws.max_row

    # --------------------------------------------------------
    # 각 ETF 컬럼의 "년도별 최소값"을 실제 데이터 셀에 표시
    # --------------------------------------------------------

    for ticker in TICKERS:

        if ticker not in header_map:
            continue

        col = header_map[ticker]

        values = []

        # --------------------------------------------
        # 해당 ETF의 데이터 수집
        # --------------------------------------------

        for row in range(2, last_data_row + 1):

            cell = ws.cell(
                row=row,
                column=col
            )

            if cell.value is None:
                continue

            try:
                value = float(cell.value)
                values.append((row, value))

            except (ValueError, TypeError):
                continue

        # 데이터가 없으면 건너뜀
        if not values:
            continue

        # --------------------------------------------
        # 해당 "년도 Sheet"에서 ETF 컬럼의 최소값
        # --------------------------------------------

        min_value = min(
            value
            for row, value in values
        )

        # --------------------------------------------
        # 최소값인 실제 셀을 빨간색 바탕으로 표시
        # --------------------------------------------

        for row, value in values:

            if value == min_value:

                cell = ws.cell(
                    row=row,
                    column=col
                )

                # 빨간색 바탕
                cell.fill = MIN_FILL

                # 빨간색 글자 + 굵게
                cell.font = MIN_FONT

                # 중앙 정렬
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

    # --------------------------------------------------------
    # 각 ETF 컬럼의 마지막 행에 최소값 표시
    # --------------------------------------------------------
    # 실제 데이터의 최소값을 다시 마지막 행에 기록합니다.
    # 날짜 컬럼에는 "최저값"을 표시합니다.
    # --------------------------------------------------------

    summary_row = last_data_row + 1

    if date_col:
        summary_date_cell = ws.cell(
            row=summary_row,
            column=date_col
        )
        summary_date_cell.value = "최저값"
        summary_date_cell.fill = SUMMARY_LABEL_FILL
        summary_date_cell.font = Font(bold=True)
        summary_date_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for ticker in TICKERS:

        if ticker not in header_map:
            continue

        col = header_map[ticker]
        values = []

        for row in range(2, last_data_row + 1):

            cell = ws.cell(
                row=row,
                column=col
            )

            if cell.value is None:
                continue

            try:
                values.append(float(cell.value))
            except (ValueError, TypeError):
                continue

        if not values:
            continue

        min_value = min(values)

        # 마지막 행에 최소값 기록
        summary_cell = ws.cell(
            row=summary_row,
            column=col
        )

        summary_cell.value = min_value
        summary_cell.fill = MIN_FILL
        summary_cell.font = MIN_FONT
        summary_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        summary_cell.number_format = "0.00"

    # --------------------------------------------------------
    # 컬럼 너비
    # --------------------------------------------------------

    if date_col:

        ws.column_dimensions[
            get_column_letter(date_col)
        ].width = 14

    for ticker in TICKERS:

        if ticker in header_map:

            col = header_map[ticker]

            ws.column_dimensions[
                get_column_letter(col)
            ].width = 12

    # --------------------------------------------------------
    # 틀 고정
    # --------------------------------------------------------

    ws.freeze_panes = "B2"

    # --------------------------------------------------------
    # 자동 필터
    # --------------------------------------------------------

    # 최저값 요약행은 필터에서 제외
    if last_data_row >= 2:

        ws.auto_filter.ref = (
            f"A1:{get_column_letter(ws.max_column)}"
            f"{last_data_row}"
        )


# ============================================================
# 모든 Sheet에 적용
# ============================================================

for ws in wb.worksheets:

    format_sheet(ws)


# ============================================================
# Excel 저장
# ============================================================

wb.save(
    OUTPUT_FILE
)


# ============================================================
# 완료
# ============================================================

print()
print("=" * 70)
print("Excel 생성 완료")
print("=" * 70)

print()
print(f"파일 : {OUTPUT_FILE}")

print()
print("전체 Sheet와 각 연도 Sheet에서")
print("ETF별 최소값 실제 셀을 빨간색으로 표시하고,")
print("각 컬럼 마지막 행에도 최소값을 기록했습니다.")

print()
print("최소값 셀 = 빨간색 바탕 + 빨간색 굵은 글씨")

